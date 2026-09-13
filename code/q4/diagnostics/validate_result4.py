"""Validate result4.xlsx against the official template and frozen Q4 arrays."""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
Q4_DIR = HERE.parent
SPEC = importlib.util.spec_from_file_location("q4_main_result4_check", Q4_DIR / "q4_main.py")
if SPEC is None or SPEC.loader is None:
    raise ImportError("无法加载q4_main.py")
q4 = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = q4
SPEC.loader.exec_module(q4)


def validate() -> dict[str, object]:
    output_path = q4.TABLES_DIR / "result4.xlsx"
    workbook = load_workbook(output_path, read_only=False, data_only=False)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"sheet名错误：{workbook.sheetnames}")
    worksheet = workbook["Sheet1"]
    expected_headers = [
        "时间\\到药材中心的距离",
        *[float(value) for value in np.round(q4.FIXED_OUTPUT_RADII_M * 100.0, 10)],
        "药材表面",
    ]
    headers = [worksheet.cell(1, column).value for column in range(1, 23)]
    if headers != expected_headers:
        raise ValueError(f"表头错误：{headers}")

    with np.load(q4.OUTPUT_DIR / "q4_solution.npz") as data:
        regular = (
            (data["time_s"] >= 60.0 - 1.0e-8)
            & (np.abs(data["time_s"] / 60.0 - np.round(data["time_s"] / 60.0)) < 1.0e-10)
        )
        expected_times = data["time_s"][regular].astype(int)
        expected = np.column_stack([data["moisture_fixed"][regular], data["moisture_surface"][regular]])
    if worksheet.max_row != len(expected_times) + 1 or worksheet.max_column != 22:
        raise ValueError(f"行列数错误：{worksheet.max_row}×{worksheet.max_column}")

    formula_count = 0
    unexpected_blank_count = 0
    expected_outside_blank_count = 0
    format_error_count = 0
    maximum_numeric_difference = 0.0
    for row_index, (time_value, expected_row) in enumerate(zip(expected_times, expected), start=2):
        if worksheet.cell(row_index, 1).value != int(time_value):
            raise ValueError(f"第{row_index}行时间错误")
        for column_index, expected_value in enumerate(expected_row, start=2):
            cell = worksheet.cell(row_index, column_index)
            formula_count += int(cell.data_type == "f")
            format_error_count += int(cell.number_format != "0.0000")
            if np.isnan(expected_value):
                expected_outside_blank_count += 1
                if cell.value is not None:
                    raise ValueError(f"体外位置应为空：{cell.coordinate}")
            else:
                if cell.value is None:
                    unexpected_blank_count += 1
                else:
                    maximum_numeric_difference = max(
                        maximum_numeric_difference,
                        abs(float(cell.value) - round(float(expected_value), 4)),
                    )
    checks = {
        "official_template": q4.record_path(q4.OFFICIAL_TEMPLATE),
        "sheet_names": workbook.sheetnames,
        "rows": worksheet.max_row,
        "columns": worksheet.max_column,
        "first_time_s": worksheet.cell(2, 1).value,
        "last_regular_time_s": worksheet.cell(worksheet.max_row, 1).value,
        "headers": headers,
        "formula_count": formula_count,
        "merged_ranges": [str(value) for value in worksheet.merged_cells.ranges],
        "unexpected_blank_count": unexpected_blank_count,
        "outside_domain_blank_count": expected_outside_blank_count,
        "number_format_error_count": format_error_count,
        "max_abs_difference_from_rounded_frozen_solution": maximum_numeric_difference,
        "all_checks_passed": bool(
            formula_count == 0
            and unexpected_blank_count == 0
            and format_error_count == 0
            and maximum_numeric_difference == 0.0
            and not worksheet.merged_cells.ranges
        ),
    }
    workbook.close()
    path = q4.OUTPUT_DIR / "result4_workbook_audit.json"
    path.write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    summary_path = q4.OUTPUT_DIR / "q4_run_summary.json"
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    summary["result_workbook"] = q4.record_path(output_path)
    summary["result_workbook_checks"] = checks
    summary["main_checks_passed"] = bool(summary["main_checks_passed"] and checks["all_checks_passed"])
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return checks


if __name__ == "__main__":
    result = validate()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["all_checks_passed"]:
        raise SystemExit(1)
