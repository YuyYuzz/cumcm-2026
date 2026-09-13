"""CUMCM 2026 Problem A, Question 2 coupled radial finite-volume solver.

The frozen mathematical specification is in model_spec_q2.md.  Q2 starts from
the prescribed initial state at t=0 and does not continue the Q1 solution.
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import numpy as np
from openpyxl import load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.optimize import brentq
from scipy.sparse import bmat, diags


Q2_ROOT = Path(__file__).resolve().parent
REPO_ROOT = Q2_ROOT.parents[1]
RESULTS_DIR = REPO_ROOT / "results" / "q2"
FIGURE_DIR = REPO_ROOT / "figures" / "final" / "q2"
VALIDATION_DIR = RESULTS_DIR
DEFAULT_XLSX = REPO_ROOT / "tables" / "q2" / "result2.xlsx"


def _find_problem_data_root() -> Path:
    """Locate original problem data without relying on a machine-specific path.

    The first layout is the formal repository convention.  The second is the
    current competition workspace layout and can be removed after migration.
    """
    for parent in Q2_ROOT.parents:
        for candidate in (
            parent / "data" / "raw" / "problems",
            parent / "CUMCM2026Problems",
        ):
            if candidate.is_dir():
                return candidate
    raise FileNotFoundError(
        "Cannot locate data/raw/problems or CUMCM2026Problems above q2_main.py"
    )


def _find_unique_input(filename: str) -> Path:
    matches = [
        path
        for path in sorted(_find_problem_data_root().rglob(filename))
        if "A题" in path.parts or "A" in path.parts
    ]
    if len(matches) != 1:
        raise FileNotFoundError(
            f"Expected exactly one {filename!r} below the problem-data root, "
            f"found {len(matches)}"
        )
    return matches[0]


ATTACHMENT_1 = _find_unique_input("附件1.xlsx")
OFFICIAL_TEMPLATE = _find_unique_input("result2.xlsx")

RADIUS_M = 0.02
LENGTH_M = 0.25
H = 25.0
HM = 8.0e-7
T_INITIAL_C = 28.0
C_INITIAL = 2.55
END_TIME_S = 10800.0
OUTPUT_TIMES_S = np.arange(0.0, END_TIME_S + 1.0, 1.0)
OUTPUT_RADII_M = np.arange(0.0, RADIUS_M + 0.0005, 0.001)
TABLE_TIMES_S = np.arange(1800.0, END_TIME_S + 1.0, 1800.0)
TABLE_RADII_M = np.array([0.0, 0.005, 0.010, 0.015, 0.020])
SNAPSHOT_TIMES_S = np.unique(
    np.r_[0.0, 1.0, 10.0, 100.0, TABLE_TIMES_S, 8880.0, 9000.0]
)


@dataclass(frozen=True)
class AmbientData:
    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray

    def temperature(self, t: float | np.ndarray) -> float | np.ndarray:
        self._check_domain(t)
        return np.interp(t, self.time_s, self.temperature_c)

    def concentration(self, t: float | np.ndarray) -> float | np.ndarray:
        self._check_domain(t)
        return np.interp(t, self.time_s, self.moisture)

    @staticmethod
    def _check_domain(t: float | np.ndarray) -> None:
        values = np.asarray(t)
        if np.any(values < -1.0e-9) or np.any(values > END_TIME_S + 1.0e-9):
            raise ValueError("Q2 ambient interpolation is defined only on 0 <= t <= 10800 s")


@dataclass(frozen=True)
class RadialGrid:
    cells: int
    faces_m: np.ndarray
    centers_m: np.ndarray
    volumes_rdr_m2: np.ndarray
    dr_m: float


@dataclass
class Q2Solution:
    grid: RadialGrid
    ambient: AmbientData
    time_s: np.ndarray
    temperature_output_c: np.ndarray
    moisture_output: np.ndarray
    temperature_surface_c: np.ndarray
    moisture_surface: np.ndarray
    ambient_temperature_c: np.ndarray
    ambient_moisture: np.ndarray
    heat_boundary_rate_w: np.ndarray
    heat_storage_cumulative_j: np.ndarray
    heat_boundary_cumulative_j: np.ndarray
    moisture_boundary_rate: np.ndarray
    moisture_inventory: np.ndarray
    snapshot_times_s: np.ndarray
    temperature_cells_snapshot_c: np.ndarray
    moisture_cells_snapshot: np.ndarray
    nfev: int
    njev: int
    nlu: int
    runtime_s: float
    tolerances: dict[str, float]


def load_q2_ambient(path: Path = ATTACHMENT_1) -> AmbientData:
    """Read Attachment 1 and retain exactly t=0,60,...,10800 s."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if not rows or len(rows[0]) < 3:
        raise ValueError("附件1 must contain at least three columns")
    headers = tuple(str(value).strip() for value in rows[0][:3])
    expected = ("时间", "温度", "水分浓度")
    if headers != expected:
        raise ValueError(f"附件1 headers are {headers}, expected {expected}")
    values = np.asarray([row[:3] for row in rows[1:] if row[0] is not None], dtype=float)
    q2 = values[values[:, 0] <= END_TIME_S]
    if q2.shape != (181, 3):
        raise ValueError(f"Expected 181 Q2 ambient rows, got {q2.shape[0]}")
    expected_times = np.arange(0.0, END_TIME_S + 60.0, 60.0)
    if not np.array_equal(q2[:, 0], expected_times):
        raise ValueError("附件1 Q2 times must be exactly 0,60,...,10800 s")
    if not np.all(np.isfinite(q2)):
        raise ValueError("附件1 Q2 data contain non-finite values")
    return AmbientData(q2[:, 0], q2[:, 1], q2[:, 2])


def make_grid(cells: int) -> RadialGrid:
    if cells < 3:
        raise ValueError("At least three finite volumes are required")
    faces = np.linspace(0.0, RADIUS_M, cells + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    volumes = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return RadialGrid(cells, faces, centers, volumes, RADIUS_M / cells)


def _return_like_input(result: np.ndarray, original: np.ndarray | float):
    return float(result) if np.asarray(original).ndim == 0 else result


def density(concentration: np.ndarray | float) -> np.ndarray | float:
    values = np.asarray(concentration, dtype=float)
    if np.any(values <= 0.0) or not np.all(np.isfinite(values)):
        raise FloatingPointError("Moisture concentration must remain finite and positive")
    result = 650.0 + 128.0 * values
    return _return_like_input(result, concentration)


def heat_capacity(concentration: np.ndarray | float) -> np.ndarray | float:
    values = np.asarray(concentration, dtype=float)
    if np.any(values <= 0.0) or not np.all(np.isfinite(values)):
        raise FloatingPointError("Moisture concentration must remain finite and positive")
    result = 1450.0 + 2736.0 * values / (values + 1.0)
    return _return_like_input(result, concentration)


def thermal_conductivity(concentration: np.ndarray | float) -> np.ndarray | float:
    values = np.asarray(concentration, dtype=float)
    if np.any(values <= 0.0) or not np.all(np.isfinite(values)):
        raise FloatingPointError("Moisture concentration must remain finite and positive")
    result = 0.21 + 0.38 * values / (values + 1.0)
    return _return_like_input(result, concentration)


def diffusion_coefficient(
    concentration: np.ndarray | float,
    temperature_c: np.ndarray | float,
) -> np.ndarray | float:
    """Appendix 3: D=2.4e-3 exp(-0.45/C) exp(-3850/(T_C+273.15))."""
    concentration_values = np.asarray(concentration, dtype=float)
    temperature_values = np.asarray(temperature_c, dtype=float)
    if np.any(concentration_values <= 0.0):
        raise FloatingPointError("Moisture concentration must remain positive")
    temperature_k = temperature_values + 273.15
    if np.any(temperature_k <= 0.0):
        raise FloatingPointError("Absolute temperature must remain positive")
    if not np.all(np.isfinite(concentration_values)) or not np.all(
        np.isfinite(temperature_values)
    ):
        raise FloatingPointError("State contains non-finite values")
    result = (
        2.4e-3
        * np.exp(-0.45 / concentration_values)
        * np.exp(-3850.0 / temperature_k)
    )
    scalar = concentration_values.ndim == 0 and temperature_values.ndim == 0
    return float(result) if scalar else result


def initial_properties() -> dict[str, float]:
    rho0 = float(density(C_INITIAL))
    cp0 = float(heat_capacity(C_INITIAL))
    k0 = float(thermal_conductivity(C_INITIAL))
    d0 = float(diffusion_coefficient(C_INITIAL, T_INITIAL_C))
    alpha0 = k0 / (rho0 * cp0)
    return {
        "rho_kg_m3": rho0,
        "cp_j_kg_k": cp0,
        "k_w_m_k": k0,
        "D_m2_s": d0,
        "alpha_m2_s": alpha0,
        "tau_temperature_s": RADIUS_M**2 / alpha0,
        "tau_moisture_s": RADIUS_M**2 / d0,
    }


def print_initial_properties() -> dict[str, float]:
    properties = initial_properties()
    print("Appendix 3 initial-property check at C=2.55, T=28 °C:")
    print(f"rho(2.55) = {properties['rho_kg_m3']:.10g} kg/m^3")
    print(f"cp(2.55)  = {properties['cp_j_kg_k']:.10g} J/(kg K)")
    print(f"k(2.55)   = {properties['k_w_m_k']:.10g} W/(m K)")
    print(f"D(2.55,28 °C) = {properties['D_m2_s']:.10e} m^2/s")
    if not (
        abs(properties["rho_kg_m3"] - 976.4) < 1.0e-9
        and 3410.0 < properties["cp_j_kg_k"] < 3420.0
        and 0.482 < properties["k_w_m_k"] < 0.484
        and 5.5e-9 < properties["D_m2_s"] < 5.8e-9
    ):
        raise RuntimeError("Appendix 3 initial-property check failed")
    return properties


def harmonic_mean(
    left: np.ndarray | float, right: np.ndarray | float
) -> np.ndarray | float:
    left_values = np.asarray(left, dtype=float)
    right_values = np.asarray(right, dtype=float)
    result = 2.0 * left_values * right_values / (left_values + right_values)
    scalar = left_values.ndim == 0 and right_values.ndim == 0
    return float(result) if scalar else result


def surface_states(
    cell_temperature_c: float,
    cell_concentration: float,
    ambient_temperature_c: float,
    ambient_concentration: float,
    dr_m: float,
) -> tuple[float, float, float, float]:
    """Reconstruct T_s, C_s and outward-coordinate fluxes k*T_r, D*C_r.

    The physical moisture branch is the first root encountered when moving
    continuously from the last cell value toward the ambient concentration.
    """
    half_width = 0.5 * dr_m
    k_cell = float(thermal_conductivity(cell_concentration))
    d_cell = float(diffusion_coefficient(cell_concentration, cell_temperature_c))

    def temperature_for_surface_concentration(surface_concentration: float) -> float:
        k_surface = float(thermal_conductivity(surface_concentration))
        k_face = float(harmonic_mean(k_cell, k_surface))
        conductance = k_face / half_width
        return float(
            (conductance * cell_temperature_c + H * ambient_temperature_c)
            / (conductance + H)
        )

    if abs(cell_concentration - ambient_concentration) <= 1.0e-14:
        surface_concentration = float(cell_concentration)
        surface_temperature = temperature_for_surface_concentration(surface_concentration)
        heat_flux = -H * (surface_temperature - ambient_temperature_c)
        return surface_temperature, surface_concentration, heat_flux, 0.0

    def residual(surface_concentration: float) -> float:
        surface_temperature = temperature_for_surface_concentration(
            surface_concentration
        )
        d_surface = float(
            diffusion_coefficient(surface_concentration, surface_temperature)
        )
        d_face = float(harmonic_mean(d_cell, d_surface))
        return (
            d_face
            * (surface_concentration - cell_concentration)
            / half_width
            + HM * (surface_concentration - ambient_concentration)
        )

    samples = np.linspace(cell_concentration, ambient_concentration, 129)
    previous_x = float(samples[0])
    previous_f = residual(previous_x)
    bracket: tuple[float, float] | None = None
    for current_x in samples[1:]:
        current_x = float(current_x)
        current_f = residual(current_x)
        if previous_f == 0.0 or previous_f * current_f <= 0.0:
            bracket = (previous_x, current_x)
            break
        previous_x, previous_f = current_x, current_f
    if bracket is None:
        raise RuntimeError("Could not bracket the continuous surface-moisture branch")
    lower, upper = sorted(bracket)
    surface_concentration = float(
        brentq(residual, lower, upper, xtol=1.0e-13, rtol=1.0e-13)
    )
    surface_temperature = temperature_for_surface_concentration(
        surface_concentration
    )
    heat_flux = -H * (surface_temperature - ambient_temperature_c)
    moisture_flux = -HM * (surface_concentration - ambient_concentration)
    return (
        surface_temperature,
        surface_concentration,
        float(heat_flux),
        float(moisture_flux),
    )


def _jacobian_sparsity(cells: int):
    tri = diags(
        [np.ones(cells - 1), np.ones(cells), np.ones(cells - 1)],
        offsets=[-1, 0, 1],
        shape=(cells, cells),
        format="csr",
    )
    return bmat([[tri, tri], [tri, tri]], format="csr")


def rhs_and_fluxes(
    t: float,
    state: np.ndarray,
    grid: RadialGrid,
    ambient: AmbientData,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, tuple[float, float]]:
    cells = grid.cells
    temperature = state[:cells]
    concentration = state[cells:]
    rho = np.asarray(density(concentration))
    cp = np.asarray(heat_capacity(concentration))
    k_cell = np.asarray(thermal_conductivity(concentration))
    diffusivity = np.asarray(diffusion_coefficient(concentration, temperature))

    heat_flux = np.zeros(cells + 1)
    moisture_flux = np.zeros(cells + 1)
    heat_flux[1:-1] = (
        harmonic_mean(k_cell[:-1], k_cell[1:])
        * np.diff(temperature)
        / grid.dr_m
    )
    moisture_flux[1:-1] = (
        harmonic_mean(diffusivity[:-1], diffusivity[1:])
        * np.diff(concentration)
        / grid.dr_m
    )
    surface_temperature, surface_concentration, heat_flux[-1], moisture_flux[-1] = (
        surface_states(
            float(temperature[-1]),
            float(concentration[-1]),
            float(ambient.temperature(t)),
            float(ambient.concentration(t)),
            grid.dr_m,
        )
    )
    temperature_divergence = (
        grid.faces_m[1:] * heat_flux[1:]
        - grid.faces_m[:-1] * heat_flux[:-1]
    ) / grid.volumes_rdr_m2
    moisture_divergence = (
        grid.faces_m[1:] * moisture_flux[1:]
        - grid.faces_m[:-1] * moisture_flux[:-1]
    ) / grid.volumes_rdr_m2
    derivative = np.r_[
        temperature_divergence / (rho * cp),
        moisture_divergence,
    ]
    return (
        derivative,
        heat_flux,
        moisture_flux,
        (surface_temperature, surface_concentration),
    )


def coupled_rhs(
    grid: RadialGrid, ambient: AmbientData
) -> Callable[[float, np.ndarray], np.ndarray]:
    def rhs(t: float, state: np.ndarray) -> np.ndarray:
        return rhs_and_fluxes(t, state, grid, ambient)[0]

    return rhs


def _axis_value(cell_values: np.ndarray, centers_m: np.ndarray) -> np.ndarray:
    r0_sq, r1_sq = centers_m[0] ** 2, centers_m[1] ** 2
    return (
        cell_values[:, 0] * r1_sq - cell_values[:, 1] * r0_sq
    ) / (r1_sq - r0_sq)


def reconstruct_to_radii(
    cell_values: np.ndarray,
    surface_values: np.ndarray,
    grid: RadialGrid,
    radii_m: np.ndarray,
) -> np.ndarray:
    axis = _axis_value(cell_values, grid.centers_m)
    reconstruction_r = np.r_[0.0, grid.centers_m, RADIUS_M]
    reconstruction_y = np.column_stack([axis, cell_values, surface_values])
    return PchipInterpolator(
        reconstruction_r, reconstruction_y, axis=1, extrapolate=False
    )(radii_m)


def _reconstruct_dense_solution(
    dense_solution,
    grid: RadialGrid,
    ambient: AmbientData,
    output_times_s: np.ndarray,
    chunk_size: int = 128,
) -> tuple[
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
    np.ndarray,
]:
    time_count = output_times_s.size
    temperature_output = np.empty((time_count, OUTPUT_RADII_M.size))
    moisture_output = np.empty_like(temperature_output)
    temperature_surface = np.empty(time_count)
    moisture_surface = np.empty(time_count)
    heat_boundary_rate = np.empty(time_count)
    heat_storage_cumulative = np.zeros(time_count)
    heat_boundary_cumulative = np.zeros(time_count)
    moisture_boundary_rate = np.empty(time_count)
    moisture_inventory = np.empty(time_count)
    geometry_factor = 2.0 * np.pi * LENGTH_M
    previous_temperature: np.ndarray | None = None
    previous_moisture: np.ndarray | None = None
    previous_time: float | None = None
    for start in range(0, time_count, chunk_size):
        stop = min(start + chunk_size, time_count)
        times = output_times_s[start:stop]
        states = dense_solution(times).T
        temperature_cells = states[:, : grid.cells]
        moisture_cells = states[:, grid.cells :]
        for local_index, (time_value, last_t, last_c) in enumerate(
            zip(times, temperature_cells[:, -1], moisture_cells[:, -1])
        ):
            surface_t, surface_c, heat_flux, moisture_flux = surface_states(
                float(last_t),
                float(last_c),
                float(ambient.temperature(time_value)),
                float(ambient.concentration(time_value)),
                grid.dr_m,
            )
            global_index = start + local_index
            temperature_surface[global_index] = surface_t
            moisture_surface[global_index] = surface_c
            heat_boundary_rate[global_index] = (
                geometry_factor * RADIUS_M * heat_flux
            )
            moisture_boundary_rate[global_index] = (
                geometry_factor * RADIUS_M * moisture_flux
            )
            moisture_inventory[global_index] = geometry_factor * np.sum(
                moisture_cells[local_index] * grid.volumes_rdr_m2
            )
            if global_index > 0:
                if (
                    previous_temperature is None
                    or previous_moisture is None
                    or previous_time is None
                ):
                    raise RuntimeError("Missing preceding state for cumulative balance")
                current_temperature = temperature_cells[local_index]
                current_moisture = moisture_cells[local_index]
                previous_capacity = np.asarray(density(previous_moisture)) * np.asarray(
                    heat_capacity(previous_moisture)
                )
                current_capacity = np.asarray(density(current_moisture)) * np.asarray(
                    heat_capacity(current_moisture)
                )
                heat_storage_cumulative[global_index] = (
                    heat_storage_cumulative[global_index - 1]
                    + geometry_factor
                    * np.sum(
                        0.5
                        * (previous_capacity + current_capacity)
                        * (current_temperature - previous_temperature)
                        * grid.volumes_rdr_m2
                    )
                )
                time_step = float(time_value - previous_time)
                heat_boundary_cumulative[global_index] = (
                    heat_boundary_cumulative[global_index - 1]
                    + 0.5
                    * (
                        heat_boundary_rate[global_index - 1]
                        + heat_boundary_rate[global_index]
                    )
                    * time_step
                )
            previous_temperature = temperature_cells[local_index]
            previous_moisture = moisture_cells[local_index]
            previous_time = float(time_value)
        temperature_output[start:stop] = reconstruct_to_radii(
            temperature_cells,
            temperature_surface[start:stop],
            grid,
            OUTPUT_RADII_M,
        )
        moisture_output[start:stop] = reconstruct_to_radii(
            moisture_cells,
            moisture_surface[start:stop],
            grid,
            OUTPUT_RADII_M,
        )
    temperature_output[0, :] = T_INITIAL_C
    moisture_output[0, :] = C_INITIAL
    temperature_surface[0] = T_INITIAL_C
    moisture_surface[0] = C_INITIAL
    initial_surface = surface_states(
        T_INITIAL_C,
        C_INITIAL,
        float(ambient.temperature(0.0)),
        float(ambient.concentration(0.0)),
        grid.dr_m,
    )
    heat_boundary_rate[0] = geometry_factor * RADIUS_M * initial_surface[2]
    moisture_boundary_rate[0] = geometry_factor * RADIUS_M * initial_surface[3]
    moisture_inventory[0] = geometry_factor * C_INITIAL * np.sum(
        grid.volumes_rdr_m2
    )
    return (
        temperature_output,
        moisture_output,
        temperature_surface,
        moisture_surface,
        heat_boundary_rate,
        heat_storage_cumulative,
        heat_boundary_cumulative,
        moisture_boundary_rate,
        moisture_inventory,
    )


def solve_q2(
    cells: int = 2561,
    rtol: float = 1.0e-8,
    atol_temperature: float = 1.0e-10,
    atol_moisture: float = 1.0e-11,
    max_step_s: float = 10.0,
    output_times_s: np.ndarray = OUTPUT_TIMES_S,
) -> Q2Solution:
    ambient = load_q2_ambient()
    grid = make_grid(cells)
    initial_state = np.r_[
        np.full(cells, T_INITIAL_C), np.full(cells, C_INITIAL)
    ]
    absolute_tolerance = np.r_[
        np.full(cells, atol_temperature),
        np.full(cells, atol_moisture),
    ]
    started = time.perf_counter()
    integration = solve_ivp(
        coupled_rhs(grid, ambient),
        t_span=(0.0, END_TIME_S),
        y0=initial_state,
        method="BDF",
        rtol=rtol,
        atol=absolute_tolerance,
        max_step=max_step_s,
        jac_sparsity=_jacobian_sparsity(cells),
        dense_output=True,
    )
    if not integration.success or integration.sol is None:
        raise RuntimeError(f"Coupled Q2 BDF failed: {integration.message}")
    output_times = np.asarray(output_times_s, dtype=float)
    (
        temperature_output,
        moisture_output,
        temperature_surface,
        moisture_surface,
        heat_boundary_rate,
        heat_storage_cumulative,
        heat_boundary_cumulative,
        moisture_boundary_rate,
        moisture_inventory,
    ) = _reconstruct_dense_solution(
        integration.sol, grid, ambient, output_times
    )
    snapshot_states = integration.sol(SNAPSHOT_TIMES_S).T
    elapsed = time.perf_counter() - started
    return Q2Solution(
        grid=grid,
        ambient=ambient,
        time_s=output_times,
        temperature_output_c=temperature_output,
        moisture_output=moisture_output,
        temperature_surface_c=temperature_surface,
        moisture_surface=moisture_surface,
        ambient_temperature_c=np.asarray(ambient.temperature(output_times)),
        ambient_moisture=np.asarray(ambient.concentration(output_times)),
        heat_boundary_rate_w=heat_boundary_rate,
        heat_storage_cumulative_j=heat_storage_cumulative,
        heat_boundary_cumulative_j=heat_boundary_cumulative,
        moisture_boundary_rate=moisture_boundary_rate,
        moisture_inventory=moisture_inventory,
        snapshot_times_s=SNAPSHOT_TIMES_S.copy(),
        temperature_cells_snapshot_c=snapshot_states[:, :cells],
        moisture_cells_snapshot=snapshot_states[:, cells:],
        nfev=int(integration.nfev),
        njev=int(integration.njev),
        nlu=int(integration.nlu),
        runtime_s=elapsed,
        tolerances={
            "rtol": rtol,
            "atol_temperature": atol_temperature,
            "atol_moisture": atol_moisture,
            "max_step_s": max_step_s,
        },
    )


def values_at(
    solution: Q2Solution, times_s: np.ndarray, radii_m: np.ndarray
) -> tuple[np.ndarray, np.ndarray]:
    time_indices = np.searchsorted(solution.time_s, times_s)
    if not np.array_equal(solution.time_s[time_indices], times_s):
        raise ValueError("Requested report times are not present")
    radius_indices = np.searchsorted(OUTPUT_RADII_M, radii_m)
    if not np.allclose(
        OUTPUT_RADII_M[radius_indices], radii_m, rtol=0.0, atol=1.0e-12
    ):
        raise ValueError("Requested report radii are not present")
    return (
        solution.temperature_output_c[np.ix_(time_indices, radius_indices)],
        solution.moisture_output[np.ix_(time_indices, radius_indices)],
    )


def save_solution_data(solution: Q2Solution) -> Path:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    path = RESULTS_DIR / "q2_solution.npz"
    np.savez_compressed(
        path,
        time_s=solution.time_s,
        output_radii_m=OUTPUT_RADII_M,
        temperature_output_c=solution.temperature_output_c,
        moisture_output=solution.moisture_output,
        temperature_surface_c=solution.temperature_surface_c,
        moisture_surface=solution.moisture_surface,
        ambient_temperature_c=solution.ambient_temperature_c,
        ambient_moisture=solution.ambient_moisture,
        heat_boundary_rate_w=solution.heat_boundary_rate_w,
        heat_storage_cumulative_j=solution.heat_storage_cumulative_j,
        heat_boundary_cumulative_j=solution.heat_boundary_cumulative_j,
        moisture_boundary_rate=solution.moisture_boundary_rate,
        moisture_inventory=solution.moisture_inventory,
        snapshot_times_s=solution.snapshot_times_s,
        cell_centers_m=solution.grid.centers_m,
        temperature_cells_snapshot_c=solution.temperature_cells_snapshot_c,
        moisture_cells_snapshot=solution.moisture_cells_snapshot,
    )
    return path


def build_result_workbook(
    solution: Q2Solution,
    output_path: Path = DEFAULT_XLSX,
) -> dict[str, object]:
    """Copy the official template and fill values only.

    The source workbook in the current problem bundle has damaged worksheet
    labels.  The two worksheet names are normalized to the names required by
    the problem statement; all template formatting and the A1 text are kept.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.resolve() == OFFICIAL_TEMPLATE.resolve():
        raise ValueError("Refusing to overwrite the official result2.xlsx template")
    shutil.copy2(OFFICIAL_TEMPLATE, output_path)
    workbook = load_workbook(output_path)
    if len(workbook.worksheets) != 2:
        raise ValueError("Official result2.xlsx must contain exactly two sheets")
    temperature_sheet, moisture_sheet = workbook.worksheets
    temperature_sheet.title = "温度"
    moisture_sheet.title = "水分浓度"
    headers = np.round(OUTPUT_RADII_M * 100.0, 10)
    temperatures = np.round(solution.temperature_output_c[1:], 4)
    moistures = np.round(solution.moisture_output[1:], 4)
    for worksheet, values in (
        (temperature_sheet, temperatures),
        (moisture_sheet, moistures),
    ):
        for column, value in enumerate(headers, start=2):
            worksheet.cell(1, column, float(value))
        for row, time_value in enumerate(solution.time_s[1:], start=2):
            worksheet.cell(row, 1, int(time_value))
            for column, value in enumerate(values[row - 2], start=2):
                worksheet.cell(row, column, float(value))
    workbook.save(output_path)
    return validate_result_workbook(output_path, solution)


def validate_result_workbook(
    path: Path, solution: Q2Solution | None = None
) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    expected_names = ["温度", "水分浓度"]
    if workbook.sheetnames != expected_names:
        raise ValueError(f"Unexpected result2.xlsx sheets: {workbook.sheetnames}")
    checks: dict[str, object] = {"sheet_names": workbook.sheetnames}
    expected_headers = np.round(OUTPUT_RADII_M * 100.0, 10)
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        if worksheet.max_row != 10801 or worksheet.max_column != 22:
            raise ValueError(
                f"{worksheet.title}: expected 10801 x 22, got "
                f"{worksheet.max_row} x {worksheet.max_column}"
            )
        rows = worksheet.iter_rows(
            min_row=1, max_row=10801, min_col=1, max_col=22, values_only=True
        )
        header_row = next(rows)
        headers = np.asarray(header_row[1:], dtype=float)
        if not np.array_equal(headers, expected_headers):
            raise ValueError(f"{worksheet.title}: radius header mismatch")
        sample_rows = {2, 101, 5401, 10801}
        expected_values = None
        if solution is not None:
            expected_values = (
                np.round(solution.temperature_output_c[1:], 4)
                if sheet_index == 0
                else np.round(solution.moisture_output[1:], 4)
            )
        formula_error_count = 0
        for excel_row, row_values in enumerate(rows, start=2):
            if row_values[0] != excel_row - 1:
                raise ValueError(
                    f"{worksheet.title}: time mismatch at row {excel_row}"
                )
            if any(
                isinstance(value, str) and value.startswith("#")
                for value in row_values[1:]
            ):
                formula_error_count += 1
            if expected_values is not None and excel_row in sample_rows:
                actual = np.asarray(row_values[1:], dtype=float)
                expected = expected_values[excel_row - 2]
                if not np.array_equal(actual, expected):
                    raise ValueError(
                        f"{worksheet.title}: values differ from solution at row {excel_row}"
                    )
        checks[worksheet.title] = {
            "rows": worksheet.max_row,
            "columns": worksheet.max_column,
            "first_time_s": 1,
            "last_time_s": 10800,
            "first_radius_cm": float(headers[0]),
            "last_radius_cm": float(headers[-1]),
        }
        checks[worksheet.title]["formula_error_count"] = formula_error_count
    workbook.close()
    checks["formula_error_count"] = sum(
        int(checks[name]["formula_error_count"]) for name in expected_names
    )
    return checks


def save_run_summary(
    solution: Q2Solution,
    data_path: Path,
    workbook_checks: dict[str, object],
) -> Path:
    table_temperature, table_moisture = values_at(
        solution, TABLE_TIMES_S, TABLE_RADII_M
    )
    properties = initial_properties()
    summary = {
        "model": "Q2 coupled conservative cell-centered radial finite volume + BDF",
        "grid_cells": solution.grid.cells,
        "runtime_s": solution.runtime_s,
        "nfev": solution.nfev,
        "njev": solution.njev,
        "nlu": solution.nlu,
        "tolerances": solution.tolerances,
        "method": "BDF",
        "input_file": ATTACHMENT_1.relative_to(REPO_ROOT).as_posix(),
        "official_template": OFFICIAL_TEMPLATE.relative_to(REPO_ROOT).as_posix(),
        "time_range_s": [0, 10800],
        "formal_output": {
            "time_s": [1, 10800, 1],
            "radius_cm": [0.0, 2.0, 0.1],
            "rounded_decimals_in_xlsx": 4,
        },
        "initial_properties": properties,
        "table_times_s": TABLE_TIMES_S.astype(int).tolist(),
        "table_radii_cm": (TABLE_RADII_M * 100.0).tolist(),
        "table_temperature_c": table_temperature.tolist(),
        "table_moisture": table_moisture.tolist(),
        "temperature_min_max_c": [
            float(np.min(solution.temperature_output_c)),
            float(np.max(solution.temperature_output_c)),
        ],
        "moisture_min_max": [
            float(np.min(solution.moisture_output)),
            float(np.max(solution.moisture_output)),
        ],
        "solution_data": data_path.relative_to(REPO_ROOT).as_posix(),
        "result_workbook": DEFAULT_XLSX.relative_to(REPO_ROOT).as_posix(),
        "result_workbook_checks": workbook_checks,
        "all_main_checks_passed": True,
    }
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    path = RESULTS_DIR / "q2_run_summary.json"
    path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cells", type=int, default=2561)
    parser.add_argument("--rtol", type=float, default=1.0e-8)
    parser.add_argument("--atol-temperature", type=float, default=1.0e-10)
    parser.add_argument("--atol-moisture", type=float, default=1.0e-11)
    parser.add_argument("--max-step", type=float, default=10.0)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--no-xlsx", action="store_true")
    args = parser.parse_args()

    print_initial_properties()
    solution = solve_q2(
        cells=args.cells,
        rtol=args.rtol,
        atol_temperature=args.atol_temperature,
        atol_moisture=args.atol_moisture,
        max_step_s=args.max_step,
    )
    data_path = save_solution_data(solution)
    workbook_checks: dict[str, object] = {"not_generated": True}
    if not args.no_xlsx:
        workbook_checks = build_result_workbook(solution, args.xlsx)
    summary_path = save_run_summary(solution, data_path, workbook_checks)
    table_temperature, table_moisture = values_at(
        solution, TABLE_TIMES_S, TABLE_RADII_M
    )
    print(
        f"Q2 solved with N={args.cells} coupled cells in {solution.runtime_s:.3f} s; "
        f"nfev={solution.nfev}, njev={solution.njev}, nlu={solution.nlu}"
    )
    print("Table 3 temperature (°C):")
    print(np.array2string(table_temperature, precision=8, suppress_small=False))
    print("Table 4 moisture (kg/kg):")
    print(np.array2string(table_moisture, precision=8, suppress_small=False))
    print(f"Solution data: {data_path}")
    print(f"Run summary: {summary_path}")
    if not args.no_xlsx:
        print(f"Official-template workbook: {args.xlsx}")
        print(json.dumps(workbook_checks, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
