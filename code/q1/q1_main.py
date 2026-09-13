"""CUMCM 2026 Problem A, Question 1: conservative radial finite-volume solver.

The mathematical specification is frozen in ``model_spec_q1.md``.  This file
contains only the Q1 implementation; it does not implement or modify Q2-Q4.
"""

from __future__ import annotations

import argparse
import json
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.optimize import brentq
from scipy.sparse import diags


ROOT = Path(__file__).resolve().parents[2]
ATTACHMENT_1 = ROOT / "data" / "raw" / "problems" / "A题" / "附件" / "附件1.xlsx"
TMP_DIR = ROOT / "tmp" / "q1"
FIGURE_DIR = ROOT / "figures" / "final" / "q1"
THREAD_OUTPUT_DIR = ROOT / "results" / "q1"
DEFAULT_XLSX = ROOT / "tables" / "q1" / "result1.xlsx"

RADIUS_M = 0.02
LENGTH_M = 0.25
RHO = 820.0
CP = 2600.0
K = 0.36
H = 25.0
HM = 8.0e-7
T_INITIAL_C = 28.0
C_INITIAL = 2.55
END_TIME_S = 1800.0
OUTPUT_TIMES_S = np.arange(0.0, END_TIME_S + 1.0, 1.0)
OUTPUT_RADII_M = np.arange(0.0, RADIUS_M + 0.0005, 0.001)
TABLE_TIMES_S = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
TABLE_RADII_M = np.array([0.0, 0.005, 0.010, 0.015, 0.020])


@dataclass(frozen=True)
class AmbientData:
    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray

    def temperature(self, t: float | np.ndarray) -> float | np.ndarray:
        self._check_domain(t)
        return np.interp(t, self.time_s, self.temperature_c)

    def concentration(self, t: float | np.ndarray) -> float | np.ndarray:
        self._check_domain(t)
        return np.interp(t, self.time_s, self.moisture)

    @staticmethod
    def _check_domain(t: float | np.ndarray) -> None:
        values = np.asarray(t)
        if np.any(values < -1.0e-9) or np.any(values > END_TIME_S + 1.0e-9):
            raise ValueError("Q1 ambient interpolation is defined only on 0 <= t <= 1800 s")


@dataclass(frozen=True)
class RadialGrid:
    cells: int
    faces_m: np.ndarray
    centers_m: np.ndarray
    volumes_rdr_m2: np.ndarray
    dr_m: float


@dataclass
class Q1Solution:
    grid: RadialGrid
    ambient: AmbientData
    time_s: np.ndarray
    temperature_cells_c: np.ndarray  # shape (time, cells)
    moisture_cells: np.ndarray  # shape (time, cells)
    temperature_surface_c: np.ndarray
    moisture_surface: np.ndarray
    temperature_output_c: np.ndarray  # shape (time, 21)
    moisture_output: np.ndarray  # shape (time, 21)
    temperature_nfev: int
    moisture_nfev: int
    runtime_s: float
    tolerances: dict[str, float]


def load_q1_ambient(path: Path = ATTACHMENT_1) -> AmbientData:
    """Read and strictly validate Attachment 1, then retain only 0--1800 s."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if not rows or len(rows[0]) < 3:
        raise ValueError("附件1 must contain at least three columns")
    headers = tuple(str(value).strip() for value in rows[0][:3])
    expected = ("时间", "温度", "水分浓度")
    if headers != expected:
        raise ValueError(f"附件1 headers are {headers}, expected {expected}")
    values = np.asarray([row[:3] for row in rows[1:] if row[0] is not None], dtype=float)
    mask = values[:, 0] <= END_TIME_S
    q1 = values[mask]
    if q1.shape != (31, 3):
        raise ValueError(f"Expected 31 Q1 ambient rows from 0 to 1800 s, got {q1.shape[0]}")
    if not np.array_equal(q1[:, 0], np.arange(0.0, 1801.0, 60.0)):
        raise ValueError("附件1 Q1 times must be exactly 0, 60, ..., 1800 s")
    if not np.all(np.isfinite(q1)):
        raise ValueError("附件1 Q1 data contain non-finite values")
    return AmbientData(q1[:, 0], q1[:, 1], q1[:, 2])


def make_grid(cells: int) -> RadialGrid:
    if cells < 3:
        raise ValueError("At least three finite volumes are required")
    faces = np.linspace(0.0, RADIUS_M, cells + 1)
    centers = 0.5 * (faces[:-1] + faces[1:])
    volumes = 0.5 * (faces[1:] ** 2 - faces[:-1] ** 2)
    return RadialGrid(cells, faces, centers, volumes, RADIUS_M / cells)


def diffusion_coefficient(concentration: np.ndarray | float) -> np.ndarray | float:
    """Appendix 2 formula D(C)=7e-9 exp(-0.89/C), in m^2/s."""
    values = np.asarray(concentration)
    if np.any(values <= 0.0) or not np.all(np.isfinite(values)):
        raise FloatingPointError("Moisture concentration must remain finite and positive")
    result = 7.0e-9 * np.exp(-0.89 / values)
    return float(result) if values.ndim == 0 else result


def harmonic_mean(left: np.ndarray | float, right: np.ndarray | float) -> np.ndarray | float:
    return 2.0 * np.asarray(left) * np.asarray(right) / (np.asarray(left) + np.asarray(right))


def heat_surface_state(cell_temperature_c: float, ambient_temperature_c: float, dr_m: float) -> tuple[float, float]:
    """Return surface temperature and F=k*dT/dr at r=R."""
    half_cell_conductance = K / (0.5 * dr_m)
    surface = (
        half_cell_conductance * cell_temperature_c + H * ambient_temperature_c
    ) / (half_cell_conductance + H)
    flux = -H * (surface - ambient_temperature_c)
    return float(surface), float(flux)


def moisture_surface_state(cell_concentration: float, ambient_concentration: float, dr_m: float) -> tuple[float, float]:
    """Return surface concentration and F=D*dC/dr at r=R.

    The last half-cell diffusivity is the harmonic mean of D at the last cell
    center and at the reconstructed surface.  The scalar nonlinear Robin
    equation is solved without adding a physical parameter.
    """
    if abs(cell_concentration - ambient_concentration) <= 1.0e-14:
        return float(cell_concentration), 0.0
    d_cell = diffusion_coefficient(cell_concentration)
    half_width = 0.5 * dr_m

    def residual(surface: float) -> float:
        d_surface = diffusion_coefficient(surface)
        d_face = harmonic_mean(d_cell, d_surface)
        return float(d_face * (surface - cell_concentration) / half_width + HM * (surface - ambient_concentration))

    # D(C) tends rapidly to zero as C approaches the low ambient value.  The
    # nonlinear half-cell equation can therefore have an additional degenerate
    # root near Ca.  The consistent finite-volume branch is the first root met
    # when moving continuously from the last cell value toward the ambient
    # value.  Selecting that branch also tends to Cs -> C_last as dr -> 0.
    samples = np.linspace(cell_concentration, ambient_concentration, 257)
    previous_x = float(samples[0])
    previous_f = residual(previous_x)
    bracket = None
    for current_x in samples[1:]:
        current_x = float(current_x)
        current_f = residual(current_x)
        if previous_f == 0.0 or previous_f * current_f <= 0.0:
            bracket = (previous_x, current_x)
            break
        previous_x, previous_f = current_x, current_f
    if bracket is None:
        raise RuntimeError("Could not bracket the physical surface-moisture branch")
    lower, upper = sorted(bracket)
    surface = brentq(residual, lower, upper, xtol=1.0e-13, rtol=1.0e-13)
    flux = -HM * (surface - ambient_concentration)
    return float(surface), float(flux)


def _jacobian_sparsity(cells: int):
    return diags(
        [np.ones(cells - 1), np.ones(cells), np.ones(cells - 1)],
        offsets=[-1, 0, 1],
        shape=(cells, cells),
        format="csr",
    )


def temperature_rhs(grid: RadialGrid, ambient: AmbientData) -> Callable[[float, np.ndarray], np.ndarray]:
    def rhs(t: float, temperature: np.ndarray) -> np.ndarray:
        flux = np.zeros(grid.cells + 1)
        flux[1:-1] = K * np.diff(temperature) / grid.dr_m
        _, flux[-1] = heat_surface_state(temperature[-1], float(ambient.temperature(t)), grid.dr_m)
        divergence = (
            grid.faces_m[1:] * flux[1:] - grid.faces_m[:-1] * flux[:-1]
        ) / grid.volumes_rdr_m2
        return divergence / (RHO * CP)

    return rhs


def moisture_rhs(grid: RadialGrid, ambient: AmbientData) -> Callable[[float, np.ndarray], np.ndarray]:
    def rhs(t: float, concentration: np.ndarray) -> np.ndarray:
        diffusivity = diffusion_coefficient(concentration)
        flux = np.zeros(grid.cells + 1)
        face_diffusivity = harmonic_mean(diffusivity[:-1], diffusivity[1:])
        flux[1:-1] = face_diffusivity * np.diff(concentration) / grid.dr_m
        _, flux[-1] = moisture_surface_state(
            concentration[-1], float(ambient.concentration(t)), grid.dr_m
        )
        return (
            grid.faces_m[1:] * flux[1:] - grid.faces_m[:-1] * flux[:-1]
        ) / grid.volumes_rdr_m2

    return rhs


def _axis_value(cell_values: np.ndarray, centers_m: np.ndarray) -> np.ndarray:
    """Linear extrapolation in r^2 using the two innermost centers."""
    r0_sq, r1_sq = centers_m[0] ** 2, centers_m[1] ** 2
    return (cell_values[:, 0] * r1_sq - cell_values[:, 1] * r0_sq) / (r1_sq - r0_sq)


def reconstruct_to_radii(
    cell_values: np.ndarray,
    surface_values: np.ndarray,
    grid: RadialGrid,
    radii_m: np.ndarray,
) -> np.ndarray:
    axis = _axis_value(cell_values, grid.centers_m)
    reconstruction_r = np.r_[0.0, grid.centers_m, RADIUS_M]
    reconstruction_y = np.column_stack([axis, cell_values, surface_values])
    return PchipInterpolator(reconstruction_r, reconstruction_y, axis=1)(radii_m)


def solve_q1(
    cells: int = 2561,
    rtol: float = 1.0e-8,
    atol_temperature: float = 1.0e-10,
    atol_moisture: float = 1.0e-11,
    max_step_s: float = 5.0,
    output_times_s: np.ndarray = OUTPUT_TIMES_S,
) -> Q1Solution:
    ambient = load_q1_ambient()
    grid = make_grid(cells)
    sparsity = _jacobian_sparsity(cells)
    started = time.perf_counter()
    common = dict(
        t_span=(0.0, END_TIME_S),
        t_eval=np.asarray(output_times_s, dtype=float),
        method="BDF",
        rtol=rtol,
        max_step=max_step_s,
        jac_sparsity=sparsity,
    )
    temperature_solution = solve_ivp(
        temperature_rhs(grid, ambient),
        y0=np.full(cells, T_INITIAL_C),
        atol=atol_temperature,
        **common,
    )
    if not temperature_solution.success:
        raise RuntimeError(f"Temperature BDF failed: {temperature_solution.message}")
    moisture_solution = solve_ivp(
        moisture_rhs(grid, ambient),
        y0=np.full(cells, C_INITIAL),
        atol=atol_moisture,
        **common,
    )
    if not moisture_solution.success:
        raise RuntimeError(f"Moisture BDF failed: {moisture_solution.message}")

    temperature_cells = temperature_solution.y.T
    moisture_cells = moisture_solution.y.T
    ambient_temperature = np.asarray(ambient.temperature(output_times_s))
    ambient_moisture = np.asarray(ambient.concentration(output_times_s))
    temperature_surface = np.array(
        [heat_surface_state(value, air, grid.dr_m)[0] for value, air in zip(temperature_cells[:, -1], ambient_temperature)]
    )
    moisture_surface = np.array(
        [moisture_surface_state(value, air, grid.dr_m)[0] for value, air in zip(moisture_cells[:, -1], ambient_moisture)]
    )
    temperature_output = reconstruct_to_radii(
        temperature_cells, temperature_surface, grid, OUTPUT_RADII_M
    )
    moisture_output = reconstruct_to_radii(
        moisture_cells, moisture_surface, grid, OUTPUT_RADII_M
    )
    # The initial field is prescribed on the closed interval, whereas the
    # Robin condition is incompatible with its zero gradient at the single
    # corner (R, 0).  Report t=0 from the prescribed initial condition exactly.
    temperature_output[0, :] = T_INITIAL_C
    moisture_output[0, :] = C_INITIAL
    elapsed = time.perf_counter() - started
    return Q1Solution(
        grid=grid,
        ambient=ambient,
        time_s=np.asarray(output_times_s),
        temperature_cells_c=temperature_cells,
        moisture_cells=moisture_cells,
        temperature_surface_c=temperature_surface,
        moisture_surface=moisture_surface,
        temperature_output_c=temperature_output,
        moisture_output=moisture_output,
        temperature_nfev=int(temperature_solution.nfev),
        moisture_nfev=int(moisture_solution.nfev),
        runtime_s=elapsed,
        tolerances={
            "rtol": rtol,
            "atol_temperature": atol_temperature,
            "atol_moisture": atol_moisture,
            "max_step_s": max_step_s,
        },
    )


def values_at(solution: Q1Solution, times_s: np.ndarray, radii_m: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    time_indices = np.searchsorted(solution.time_s, times_s)
    if not np.array_equal(solution.time_s[time_indices], times_s):
        raise ValueError("Requested report times are not present in the solution output")
    radius_indices = np.searchsorted(OUTPUT_RADII_M, radii_m)
    if not np.allclose(OUTPUT_RADII_M[radius_indices], radii_m, rtol=0.0, atol=1.0e-12):
        raise ValueError("Requested report radii are not present in the output grid")
    return (
        solution.temperature_output_c[np.ix_(time_indices, radius_indices)],
        solution.moisture_output[np.ix_(time_indices, radius_indices)],
    )


def save_solution_data(solution: Q1Solution) -> Path:
    THREAD_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = THREAD_OUTPUT_DIR / f"q1_solution_n{solution.grid.cells}.npz"
    np.savez_compressed(
        path,
        time_s=solution.time_s,
        cell_centers_m=solution.grid.centers_m,
        output_radii_m=OUTPUT_RADII_M,
        temperature_cells_c=solution.temperature_cells_c,
        moisture_cells=solution.moisture_cells,
        temperature_surface_c=solution.temperature_surface_c,
        moisture_surface=solution.moisture_surface,
        temperature_output_c=solution.temperature_output_c,
        moisture_output=solution.moisture_output,
    )
    return path


def save_figures(solution: Q1Solution) -> list[Path]:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    dense_r = np.linspace(0.0, RADIUS_M, 201)
    dense_temperature = reconstruct_to_radii(
        solution.temperature_cells_c, solution.temperature_surface_c, solution.grid, dense_r
    )
    dense_moisture = reconstruct_to_radii(
        solution.moisture_cells, solution.moisture_surface, solution.grid, dense_r
    )
    profile_times = [100, 600, 1200, 1800]
    indices = [int(np.searchsorted(solution.time_s, value)) for value in profile_times]
    colors = ["#2563EB", "#0F9D8A", "#F59E0B", "#DC2626"]

    plt.rcParams.update({"font.size": 10, "axes.grid": True, "grid.alpha": 0.25})
    fig, ax = plt.subplots(figsize=(7.2, 4.8))
    for index, label, color in zip(indices, profile_times, colors):
        ax.plot(dense_r * 100.0, dense_temperature[index], lw=2.0, color=color, label=f"{label} s")
    ax.set_xlabel("Distance from cylinder axis (cm)")
    ax.set_ylabel("Temperature (°C)")
    ax.set_title("Q1 radial temperature profiles")
    ax.legend(frameon=False)
    fig.tight_layout()
    temperature_path = FIGURE_DIR / "q1_temperature_profiles.png"
    fig.savefig(temperature_path, dpi=220)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(7.2, 4.8))
    for index, label, color in zip(indices, profile_times, colors):
        ax.plot(dense_r * 100.0, dense_moisture[index], lw=2.0, color=color, label=f"{label} s")
    ax.set_xlabel("Distance from cylinder axis (cm)")
    ax.set_ylabel("Moisture concentration (kg/kg)")
    ax.set_title("Q1 radial moisture profiles")
    ax.legend(frameon=False)
    fig.tight_layout()
    moisture_path = FIGURE_DIR / "q1_moisture_profiles.png"
    fig.savefig(moisture_path, dpi=220)
    plt.close(fig)

    fig, axes = plt.subplots(2, 1, figsize=(7.6, 7.0), sharex=True)
    axes[0].plot(solution.time_s, solution.temperature_output_c[:, 0], label="Axis", lw=1.8)
    axes[0].plot(solution.time_s, solution.temperature_output_c[:, -1], label="Surface", lw=1.8)
    axes[0].plot(solution.time_s, solution.ambient.temperature(solution.time_s), label="Drying room", lw=1.3, ls="--")
    axes[0].set_ylabel("Temperature (°C)")
    axes[0].legend(frameon=False, ncol=3)
    axes[1].plot(solution.time_s, solution.moisture_output[:, 0], label="Axis", lw=1.8)
    axes[1].plot(solution.time_s, solution.moisture_output[:, -1], label="Surface", lw=1.8)
    axes[1].plot(solution.time_s, solution.ambient.concentration(solution.time_s), label="Drying room", lw=1.3, ls="--")
    axes[1].set_xlabel("Time (s)")
    axes[1].set_ylabel("Moisture concentration (kg/kg)")
    axes[1].legend(frameon=False, ncol=3)
    fig.suptitle("Q1 axis, surface, and drying-room histories")
    fig.tight_layout()
    histories_path = FIGURE_DIR / "q1_center_surface_timeseries.png"
    fig.savefig(histories_path, dpi=220)
    plt.close(fig)
    return [temperature_path, moisture_path, histories_path]


def write_workbook_payload(solution: Q1Solution) -> Path:
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    payload_path = TMP_DIR / "q1_workbook_payload.json"
    payload = {
        "time_s": [int(value) for value in solution.time_s],
        "radius_cm": np.round(OUTPUT_RADII_M * 100.0, 10).tolist(),
        "temperature_c": np.round(solution.temperature_output_c, 4).tolist(),
        "moisture": np.round(solution.moisture_output, 4).tolist(),
        "source": "data/raw/problems/A题/附件/附件1.xlsx; rows t=0..1800 s only",
        "model_grid_cells": solution.grid.cells,
    }
    payload_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return payload_path


def build_result_workbook(payload_path: Path, output_path: Path) -> None:
    payload = json.loads(payload_path.read_text(encoding="utf-8"))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.active.title = "温度"
    # Official result1.xlsx begins at t=1 s; retain t=0 only in the NPZ state.
    times = payload["time_s"][1:]
    for name, values in (("温度", payload["temperature_c"][1:]), ("水分浓度", payload["moisture"][1:])):
        worksheet = workbook[name] if name in workbook.sheetnames else workbook.create_sheet(name)
        worksheet.delete_rows(1, worksheet.max_row)
        worksheet.append(["时间\\到药材中心的距离", *payload["radius_cm"]])
        for t, row in zip(times, values):
            worksheet.append([t, *row])
    workbook.save(output_path)


def save_run_summary(solution: Q1Solution, workbook_path: Path, data_path: Path, figures: list[Path]) -> Path:
    table_t, table_c = values_at(solution, TABLE_TIMES_S, TABLE_RADII_M)
    summary = {
        "model": "Q1 conservative cell-centered radial finite volume + BDF",
        "grid_cells": solution.grid.cells,
        "runtime_s": solution.runtime_s,
        "temperature_nfev": solution.temperature_nfev,
        "moisture_nfev": solution.moisture_nfev,
        "tolerances": solution.tolerances,
        "table_times_s": TABLE_TIMES_S.astype(int).tolist(),
        "table_radii_cm": (TABLE_RADII_M * 100.0).tolist(),
        "table_temperature_c": table_t.tolist(),
        "table_moisture": table_c.tolist(),
        "temperature_min_max_c": [float(np.min(solution.temperature_output_c)), float(np.max(solution.temperature_output_c))],
        "moisture_min_max": [float(np.min(solution.moisture_output)), float(np.max(solution.moisture_output))],
        "workbook": str(workbook_path),
        "solution_data": str(data_path),
        "figures": [str(path) for path in figures],
    }
    THREAD_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    path = THREAD_OUTPUT_DIR / "q1_run_summary.json"
    path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cells", type=int, default=2561)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--no-xlsx", action="store_true")
    args = parser.parse_args()

    solution = solve_q1(cells=args.cells)
    data_path = save_solution_data(solution)
    figure_paths = save_figures(solution)
    workbook_path = args.xlsx.resolve()
    if not args.no_xlsx:
        payload_path = write_workbook_payload(solution)
        build_result_workbook(payload_path, workbook_path)
    summary_path = save_run_summary(solution, workbook_path, data_path, figure_paths)
    table_t, table_c = values_at(solution, TABLE_TIMES_S, TABLE_RADII_M)
    print(f"Q1 solved with N={args.cells} cells in {solution.runtime_s:.3f} s")
    print("Temperature table (°C):")
    print(np.array2string(table_t, precision=6, suppress_small=False))
    print("Moisture table (kg/kg):")
    print(np.array2string(table_c, precision=6, suppress_small=False))
    print(f"Workbook: {workbook_path}")
    print(f"Run summary: {summary_path}")


if __name__ == "__main__":
    main()
