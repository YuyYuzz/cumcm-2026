# 文件名：q3_main.py
# 用途：该代码负责第三问全域含水率阈值模型的求解、烘干终止时间计算及结果表格输出。

from __future__ import annotations

import argparse
import copy
import importlib.util
import json
import math
import shutil
import sys
import time
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from scipy.integrate import solve_ivp
from scipy.optimize import brentq


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
CODE_DIR = SCRIPT_DIR
RESULTS_DIR = SCRIPT_DIR
TABLES_DIR = SCRIPT_DIR


def _find_problem_data_root() -> Path:
    """定位题目官方输入所在的相对目录。"""
    for parent in (SCRIPT_DIR, *SCRIPT_DIR.parents):
        for candidate in (
            parent / "data" / "raw" / "problems",
            parent / "CUMCM2026Problems",
        ):
            if candidate.is_dir():
                return candidate
    raise FileNotFoundError("无法定位题目官方输入目录")


PROBLEM_DATA_ROOT = _find_problem_data_root()
ATTACHMENT_1 = PROBLEM_DATA_ROOT / "A题" / "附件" / "附件1.xlsx"
OFFICIAL_TEMPLATE = PROBLEM_DATA_ROOT / "A题" / "附件" / "附件3" / "result3.xlsx"
Q2_MODULE_PATH = PROJECT_ROOT / "Q2" / "q2_main.py"
Q2_SOLUTION_PATH = PROJECT_ROOT / "Q2" / "results" / "q2_solution.npz"
DEFAULT_XLSX = TABLES_DIR / "result3.xlsx"


def record_path(path: Path) -> str:
    """以相对路径记录输入输出位置。"""
    resolved = path.resolve()
    for root in (PROJECT_ROOT.resolve(), PROJECT_ROOT.parent.resolve()):
        try:
            return resolved.relative_to(root).as_posix()
        except ValueError:
            continue
    return path.name


def _load_q2_module():
    spec = importlib.util.spec_from_file_location("cumcm_q2", Q2_MODULE_PATH)
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载问题二模块：{Q2_MODULE_PATH}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


q2 = _load_q2_module()

RADIUS_M = q2.RADIUS_M
LENGTH_M = q2.LENGTH_M
H = q2.H
HM = q2.HM
T_INITIAL_C = q2.T_INITIAL_C
C_INITIAL = q2.C_INITIAL
AMBIENT_DATA_END_S = 14400.0
INTEGRATION_UPPER_S = 96.0 * 3600.0
THRESHOLD = 0.15
OUTPUT_INTERVAL_S = 60.0
OUTPUT_RADII_M = q2.OUTPUT_RADII_M.copy()
TABLE_TIMES_S = np.arange(6.0, 55.0, 6.0) * 3600.0
TABLE_RADII_M = np.array([0.0, 0.005, 0.010, 0.015, 0.020])
SNAPSHOT_BASE_S = np.unique(
    np.r_[0.0, 1.0, 10.0, 100.0, AMBIENT_DATA_END_S, TABLE_TIMES_S]
)


@dataclass(frozen=True)
class AmbientData:
    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    post_temperature_c: float
    post_moisture: float
    post_from_inclusive: bool = False

    @staticmethod
    def _return(result: np.ndarray, original: float | np.ndarray):
        return float(result) if np.asarray(original).ndim == 0 else result

    def temperature(self, t: float | np.ndarray) -> float | np.ndarray:
        values = np.asarray(t, dtype=float)
        if np.any(values < -1.0e-9):
            raise ValueError("环境时间不得小于 0")
        clipped = np.minimum(values, AMBIENT_DATA_END_S)
        result = np.interp(clipped, self.time_s, self.temperature_c)
        after = values >= AMBIENT_DATA_END_S if self.post_from_inclusive else values > AMBIENT_DATA_END_S
        result = np.where(after, self.post_temperature_c, result)
        return self._return(np.asarray(result), t)

    def concentration(self, t: float | np.ndarray) -> float | np.ndarray:
        values = np.asarray(t, dtype=float)
        if np.any(values < -1.0e-9):
            raise ValueError("环境时间不得小于 0")
        clipped = np.minimum(values, AMBIENT_DATA_END_S)
        result = np.interp(clipped, self.time_s, self.moisture)
        after = values >= AMBIENT_DATA_END_S if self.post_from_inclusive else values > AMBIENT_DATA_END_S
        result = np.where(after, self.post_moisture, result)
        return self._return(np.asarray(result), t)


@dataclass
class Q3Solution:
    grid: object
    ambient: AmbientData
    time_s: np.ndarray
    temperature_output_c: np.ndarray
    moisture_output: np.ndarray
    temperature_surface_c: np.ndarray
    moisture_surface: np.ndarray
    ambient_temperature_c: np.ndarray
    ambient_moisture: np.ndarray
    heat_boundary_rate_w: np.ndarray
    heat_storage_cumulative_j: np.ndarray
    heat_boundary_cumulative_j: np.ndarray
    moisture_boundary_rate: np.ndarray
    moisture_inventory: np.ndarray
    cmax: np.ndarray
    cmax_radius_m: np.ndarray
    snapshot_times_s: np.ndarray
    temperature_cells_snapshot_c: np.ndarray
    moisture_cells_snapshot: np.ndarray
    terminal_radii_m: np.ndarray
    terminal_moisture: np.ndarray
    end_time_s: float
    end_max_moisture: float
    end_max_radius_m: float
    nfev: int
    njev: int
    nlu: int
    runtime_s: float
    tolerances: dict[str, float]


def load_q3_ambient(
    path: Path = ATTACHMENT_1,
    post_moisture: float | None = None,
    post_from_inclusive: bool = False,
) -> AmbientData:
    """读取附件1完整 0,60,...,14400 s 数据，末端以后保持常值。"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    if not rows or len(rows[0]) < 3:
        raise ValueError("附件1至少需要三列")
    values = np.asarray([row[:3] for row in rows[1:] if row[0] is not None], dtype=float)
    expected_times = np.arange(0.0, AMBIENT_DATA_END_S + 60.0, 60.0)
    if values.shape != (241, 3):
        raise ValueError(f"附件1应有241条环境记录，实际为{values.shape[0]}条")
    if not np.array_equal(values[:, 0], expected_times):
        raise ValueError("附件1时间必须严格为0,60,...,14400 s")
    if not np.all(np.isfinite(values)):
        raise ValueError("附件1包含非有限数值")
    last_temperature = float(values[-1, 1])
    last_moisture = float(values[-1, 2])
    if abs(last_temperature - 50.165) > 1.0e-12 or abs(last_moisture - 0.04986) > 1.0e-12:
        raise ValueError(
            "附件1末值与题目核验值不一致："
            f"Ta={last_temperature}, Ca={last_moisture}"
        )
    return AmbientData(
        time_s=values[:, 0],
        temperature_c=values[:, 1],
        moisture=values[:, 2],
        post_temperature_c=last_temperature,
        post_moisture=last_moisture if post_moisture is None else float(post_moisture),
        post_from_inclusive=post_from_inclusive,
    )


class PiecewiseDenseSolution:
    """把14400 s前后两段BDF稠密输出组合成统一可调用对象。"""

    def __init__(self, first, second, switch_s: float):
        self.first = first
        self.second = second
        self.switch_s = switch_s

    def __call__(self, t):
        values = np.asarray(t, dtype=float)
        scalar = values.ndim == 0
        flat = values.reshape(-1)
        sample = self.first(float(min(flat[0], self.switch_s)))
        result = np.empty((sample.size, flat.size))
        first_mask = flat <= self.switch_s
        if np.any(first_mask):
            result[:, first_mask] = self.first(flat[first_mask])
        if np.any(~first_mask):
            result[:, ~first_mask] = self.second(flat[~first_mask])
        return result[:, 0] if scalar else result.reshape((sample.size,) + values.shape)


def axis_value_from_state(concentration: np.ndarray, grid) -> float:
    r0_sq = grid.centers_m[0] ** 2
    r1_sq = grid.centers_m[1] ** 2
    return float(
        (concentration[0] * r1_sq - concentration[1] * r0_sq)
        / (r1_sq - r0_sq)
    )


def surface_states(
    cell_temperature_c: float,
    cell_concentration: float,
    ambient_temperature_c: float,
    ambient_concentration: float,
    dr_m: float,
) -> tuple[float, float, float, float]:
    """与问题二相同的表面物理解，使用完整物理区间求根。"""
    half_width = 0.5 * dr_m
    k_cell = 0.21 + 0.38 * cell_concentration / (cell_concentration + 1.0)
    d_cell = (
        2.4e-3
        * math.exp(-0.45 / cell_concentration)
        * math.exp(-3850.0 / (cell_temperature_c + 273.15))
    )

    def temperature_for_surface_concentration(surface_concentration: float) -> float:
        k_surface = 0.21 + 0.38 * surface_concentration / (surface_concentration + 1.0)
        k_face = 2.0 * k_cell * k_surface / (k_cell + k_surface)
        conductance = k_face / half_width
        return float(
            conductance * cell_temperature_c + H * ambient_temperature_c
        ) / float(conductance + H)

    if abs(cell_concentration - ambient_concentration) <= 1.0e-14:
        surface_temperature = float(
            temperature_for_surface_concentration(cell_concentration)
        )
        return (
            surface_temperature,
            float(cell_concentration),
            float(-H * (surface_temperature - ambient_temperature_c)),
            0.0,
        )

    def residual(surface_concentration: float) -> float:
        surface_temperature = temperature_for_surface_concentration(
            surface_concentration
        )
        d_surface = (
            2.4e-3
            * math.exp(-0.45 / surface_concentration)
            * math.exp(-3850.0 / (surface_temperature + 273.15))
        )
        d_face = 2.0 * d_cell * d_surface / (d_cell + d_surface)
        return float(
            d_face
            * (surface_concentration - cell_concentration)
            / half_width
            + HM * (surface_concentration - ambient_concentration)
        )

    samples = np.linspace(cell_concentration, ambient_concentration, 129)
    k_surface_samples = 0.21 + 0.38 * samples / (samples + 1.0)
    k_face_samples = 2.0 * k_cell * k_surface_samples / (k_cell + k_surface_samples)
    conductance_samples = k_face_samples / half_width
    temperature_samples = (
        conductance_samples * cell_temperature_c + H * ambient_temperature_c
    ) / (conductance_samples + H)
    d_surface_samples = (
        2.4e-3
        * np.exp(-0.45 / samples)
        * np.exp(-3850.0 / (temperature_samples + 273.15))
    )
    d_face_samples = 2.0 * d_cell * d_surface_samples / (d_cell + d_surface_samples)
    residual_samples = (
        d_face_samples * (samples - cell_concentration) / half_width
        + HM * (samples - ambient_concentration)
    )
    crossings = np.flatnonzero(
        (residual_samples[:-1] == 0.0)
        | (residual_samples[:-1] * residual_samples[1:] <= 0.0)
    )
    if crossings.size == 0:
        raise RuntimeError("无法括住连续的表面水分物理解")
    # 与问题二一致：从最外控制体状态向环境状态扫描，取遇到的首个连续物理解。
    first = int(crossings[0])
    lower, upper = sorted((float(samples[first]), float(samples[first + 1])))
    surface_concentration = float(
        brentq(residual, lower, upper, xtol=1.0e-13, rtol=1.0e-13)
    )
    surface_temperature = float(
        temperature_for_surface_concentration(surface_concentration)
    )
    return (
        surface_temperature,
        surface_concentration,
        float(-H * (surface_temperature - ambient_temperature_c)),
        float(-HM * (surface_concentration - ambient_concentration)),
    )


def rhs_and_fluxes(t: float, state: np.ndarray, grid, ambient: AmbientData):
    cells = grid.cells
    temperature = state[:cells]
    concentration = state[cells:]
    rho = np.asarray(q2.density(concentration))
    cp = np.asarray(q2.heat_capacity(concentration))
    k_cell = np.asarray(q2.thermal_conductivity(concentration))
    diffusivity = np.asarray(q2.diffusion_coefficient(concentration, temperature))
    heat_flux = np.zeros(cells + 1)
    moisture_flux = np.zeros(cells + 1)
    heat_flux[1:-1] = (
        q2.harmonic_mean(k_cell[:-1], k_cell[1:])
        * np.diff(temperature)
        / grid.dr_m
    )
    moisture_flux[1:-1] = (
        q2.harmonic_mean(diffusivity[:-1], diffusivity[1:])
        * np.diff(concentration)
        / grid.dr_m
    )
    surface_temperature, surface_concentration, heat_flux[-1], moisture_flux[-1] = (
        surface_states(
            float(temperature[-1]),
            float(concentration[-1]),
            float(ambient.temperature(t)),
            float(ambient.concentration(t)),
            grid.dr_m,
        )
    )
    temperature_divergence = (
        grid.faces_m[1:] * heat_flux[1:]
        - grid.faces_m[:-1] * heat_flux[:-1]
    ) / grid.volumes_rdr_m2
    moisture_divergence = (
        grid.faces_m[1:] * moisture_flux[1:]
        - grid.faces_m[:-1] * moisture_flux[:-1]
    ) / grid.volumes_rdr_m2
    derivative = np.r_[
        temperature_divergence / (rho * cp),
        moisture_divergence,
    ]
    return derivative, heat_flux, moisture_flux, (
        surface_temperature,
        surface_concentration,
    )


def coupled_rhs(grid, ambient: AmbientData):
    def rhs(t: float, state: np.ndarray) -> np.ndarray:
        return rhs_and_fluxes(t, state, grid, ambient)[0]

    return rhs


def cmax_and_location(
    t: float,
    state: np.ndarray,
    grid,
    ambient: AmbientData,
) -> tuple[float, float]:
    cells = grid.cells
    temperature = state[:cells]
    concentration = state[cells:]
    axis = axis_value_from_state(concentration, grid)
    _, surface, _, _ = surface_states(
        float(temperature[-1]),
        float(concentration[-1]),
        float(ambient.temperature(t)),
        float(ambient.concentration(t)),
        grid.dr_m,
    )
    candidate_values = np.r_[axis, concentration, surface]
    candidate_radii = np.r_[0.0, grid.centers_m, RADIUS_M]
    index = int(np.argmax(candidate_values))
    return float(candidate_values[index]), float(candidate_radii[index])


def make_threshold_event(grid, ambient: AmbientData):
    def event(t: float, state: np.ndarray) -> float:
        return cmax_and_location(t, state, grid, ambient)[0] - THRESHOLD

    event.terminal = True
    event.direction = -1.0
    return event


def _output_times(end_time_s: float, requested: np.ndarray | None) -> np.ndarray:
    if requested is None:
        last_regular = np.floor(end_time_s / OUTPUT_INTERVAL_S) * OUTPUT_INTERVAL_S
        values = np.arange(0.0, last_regular + 0.5 * OUTPUT_INTERVAL_S, OUTPUT_INTERVAL_S)
    else:
        values = np.asarray(requested, dtype=float)
        if values.ndim != 1 or values.size == 0 or values[0] != 0.0:
            raise ValueError("输出时刻必须是一维数组且首项为0")
        values = values[values <= end_time_s + 1.0e-8]
    if abs(values[-1] - end_time_s) > 1.0e-7:
        values = np.r_[values, end_time_s]
    else:
        values[-1] = end_time_s
    return np.unique(values)


def _max_history(dense_solution, times_s: np.ndarray, grid, ambient: AmbientData):
    maximum = np.empty(times_s.size)
    radius = np.empty(times_s.size)
    for start in range(0, times_s.size, 64):
        stop = min(start + 64, times_s.size)
        states = dense_solution(times_s[start:stop]).T
        for local, (time_value, state) in enumerate(zip(times_s[start:stop], states)):
            maximum[start + local], radius[start + local] = cmax_and_location(
                float(time_value), state, grid, ambient
            )
    return maximum, radius


def _reconstruct_dense_solution(dense_solution, grid, ambient, output_times_s):
    time_count = output_times_s.size
    temperature_output = np.empty((time_count, OUTPUT_RADII_M.size))
    moisture_output = np.empty_like(temperature_output)
    temperature_surface = np.empty(time_count)
    moisture_surface = np.empty(time_count)
    heat_boundary_rate = np.empty(time_count)
    heat_storage_cumulative = np.zeros(time_count)
    heat_boundary_cumulative = np.zeros(time_count)
    moisture_boundary_rate = np.empty(time_count)
    moisture_inventory = np.empty(time_count)
    geometry_factor = 2.0 * np.pi * LENGTH_M
    previous_temperature = None
    previous_moisture = None
    previous_time = None
    for start in range(0, time_count, 128):
        stop = min(start + 128, time_count)
        times = output_times_s[start:stop]
        states = dense_solution(times).T
        temperature_cells = states[:, : grid.cells]
        moisture_cells = states[:, grid.cells :]
        for local_index, (time_value, last_t, last_c) in enumerate(
            zip(times, temperature_cells[:, -1], moisture_cells[:, -1])
        ):
            surface_t, surface_c, heat_flux, moisture_flux = surface_states(
                float(last_t),
                float(last_c),
                float(ambient.temperature(time_value)),
                float(ambient.concentration(time_value)),
                grid.dr_m,
            )
            global_index = start + local_index
            temperature_surface[global_index] = surface_t
            moisture_surface[global_index] = surface_c
            heat_boundary_rate[global_index] = geometry_factor * RADIUS_M * heat_flux
            moisture_boundary_rate[global_index] = geometry_factor * RADIUS_M * moisture_flux
            moisture_inventory[global_index] = geometry_factor * np.sum(
                moisture_cells[local_index] * grid.volumes_rdr_m2
            )
            if global_index > 0:
                previous_capacity = np.asarray(q2.density(previous_moisture)) * np.asarray(
                    q2.heat_capacity(previous_moisture)
                )
                current_capacity = np.asarray(q2.density(moisture_cells[local_index])) * np.asarray(
                    q2.heat_capacity(moisture_cells[local_index])
                )
                heat_storage_cumulative[global_index] = (
                    heat_storage_cumulative[global_index - 1]
                    + geometry_factor
                    * np.sum(
                        0.5
                        * (previous_capacity + current_capacity)
                        * (temperature_cells[local_index] - previous_temperature)
                        * grid.volumes_rdr_m2
                    )
                )
                time_step = float(time_value - previous_time)
                heat_boundary_cumulative[global_index] = (
                    heat_boundary_cumulative[global_index - 1]
                    + 0.5
                    * (heat_boundary_rate[global_index - 1] + heat_boundary_rate[global_index])
                    * time_step
                )
            previous_temperature = temperature_cells[local_index]
            previous_moisture = moisture_cells[local_index]
            previous_time = float(time_value)
        temperature_output[start:stop] = q2.reconstruct_to_radii(
            temperature_cells, temperature_surface[start:stop], grid, OUTPUT_RADII_M
        )
        moisture_output[start:stop] = q2.reconstruct_to_radii(
            moisture_cells, moisture_surface[start:stop], grid, OUTPUT_RADII_M
        )
    temperature_output[0, :] = T_INITIAL_C
    moisture_output[0, :] = C_INITIAL
    temperature_surface[0] = T_INITIAL_C
    moisture_surface[0] = C_INITIAL
    initial_surface = surface_states(
        T_INITIAL_C,
        C_INITIAL,
        float(ambient.temperature(0.0)),
        float(ambient.concentration(0.0)),
        grid.dr_m,
    )
    heat_boundary_rate[0] = geometry_factor * RADIUS_M * initial_surface[2]
    moisture_boundary_rate[0] = geometry_factor * RADIUS_M * initial_surface[3]
    moisture_inventory[0] = geometry_factor * C_INITIAL * np.sum(grid.volumes_rdr_m2)
    return (
        temperature_output,
        moisture_output,
        temperature_surface,
        moisture_surface,
        heat_boundary_rate,
        heat_storage_cumulative,
        heat_boundary_cumulative,
        moisture_boundary_rate,
        moisture_inventory,
    )


def solve_q3(
    cells: int = 2561,
    rtol: float = 1.0e-8,
    atol_temperature: float = 1.0e-10,
    atol_moisture: float = 1.0e-11,
    max_step_s: float = 60.0,
    integration_upper_s: float = INTEGRATION_UPPER_S,
    terminal_event: bool = True,
    output_times_s: np.ndarray | None = None,
    post_moisture: float | None = None,
) -> Q3Solution:
    ambient = load_q3_ambient(post_moisture=post_moisture)
    grid = q2.make_grid(cells)
    initial_state = np.r_[
        np.full(cells, T_INITIAL_C),
        np.full(cells, C_INITIAL),
    ]
    absolute_tolerance = np.r_[
        np.full(cells, atol_temperature),
        np.full(cells, atol_moisture),
    ]
    started = time.perf_counter()
    common_options = dict(
        method="BDF",
        rtol=rtol,
        atol=absolute_tolerance,
        max_step=max_step_s,
        jac_sparsity=q2._jacobian_sparsity(cells),
        dense_output=True,
    )
    has_boundary_step = (
        post_moisture is not None
        and abs(float(post_moisture) - float(ambient.moisture[-1])) > 1.0e-14
        and integration_upper_s > AMBIENT_DATA_END_S
    )
    if has_boundary_step:
        first_ambient = load_q3_ambient()
        first = solve_ivp(
            coupled_rhs(grid, first_ambient),
            t_span=(0.0, AMBIENT_DATA_END_S),
            y0=initial_state,
            events=None,
            **common_options,
        )
        if not first.success or first.sol is None:
            raise RuntimeError(f"Q3敏感性前段BDF积分失败：{first.message}")
        second_ambient = load_q3_ambient(
            post_moisture=post_moisture, post_from_inclusive=True
        )
        second_event = make_threshold_event(grid, second_ambient) if terminal_event else None
        second = solve_ivp(
            coupled_rhs(grid, second_ambient),
            t_span=(AMBIENT_DATA_END_S, float(integration_upper_s)),
            y0=first.y[:, -1],
            events=second_event,
            **common_options,
        )
        if not second.success or second.sol is None:
            raise RuntimeError(
                f"Q3敏感性后段BDF积分失败于t={second.t[-1]:.12g} s：{second.message}"
            )
        post_dense = second.sol
        event_times = second.t_events
        nfev = first.nfev + second.nfev
        njev = first.njev + second.njev
        nlu = first.nlu + second.nlu
        dense_solution = PiecewiseDenseSolution(first.sol, post_dense, AMBIENT_DATA_END_S)
    else:
        events = make_threshold_event(grid, ambient) if terminal_event else None
        integration = solve_ivp(
            coupled_rhs(grid, ambient),
            t_span=(0.0, float(integration_upper_s)),
            y0=initial_state,
            events=events,
            **common_options,
        )
        if not integration.success or integration.sol is None:
            raise RuntimeError(f"Q3 BDF积分失败：{integration.message}")
        dense_solution = integration.sol
        event_times = integration.t_events
        nfev, njev, nlu = integration.nfev, integration.njev, integration.nlu
    if terminal_event:
        if event_times is None or event_times[0].size != 1:
            raise RuntimeError(
                f"在{integration_upper_s / 3600:.1f} h上界内未检测到全域阈值事件"
            )
        end_time_s = float(event_times[0][0])
    else:
        end_time_s = float(integration_upper_s)
    times = _output_times(end_time_s, output_times_s)
    (
        temperature_output,
        moisture_output,
        temperature_surface,
        moisture_surface,
        heat_boundary_rate,
        heat_storage_cumulative,
        heat_boundary_cumulative,
        moisture_boundary_rate,
        moisture_inventory,
    ) = _reconstruct_dense_solution(dense_solution, grid, ambient, times)
    cmax, cmax_radius = _max_history(dense_solution, times, grid, ambient)

    snapshot_times = SNAPSHOT_BASE_S[SNAPSHOT_BASE_S <= end_time_s + 1.0e-8]
    if abs(snapshot_times[-1] - end_time_s) > 1.0e-7:
        snapshot_times = np.r_[snapshot_times, end_time_s]
    snapshot_states = dense_solution(snapshot_times).T

    terminal_state = dense_solution(end_time_s)
    end_max, end_radius = cmax_and_location(end_time_s, terminal_state, grid, ambient)
    terminal_radii = np.linspace(0.0, RADIUS_M, 10001)
    terminal_surface = surface_states(
        float(terminal_state[cells - 1]),
        float(terminal_state[-1]),
        float(ambient.temperature(end_time_s)),
        float(ambient.concentration(end_time_s)),
        grid.dr_m,
    )[1]
    terminal_moisture = q2.reconstruct_to_radii(
        terminal_state[cells:][None, :],
        np.array([terminal_surface]),
        grid,
        terminal_radii,
    )[0]
    fine_index = int(np.argmax(terminal_moisture))
    if terminal_moisture[fine_index] > end_max + 1.0e-10:
        end_max = float(terminal_moisture[fine_index])
        end_radius = float(terminal_radii[fine_index])

    elapsed = time.perf_counter() - started
    return Q3Solution(
        grid=grid,
        ambient=ambient,
        time_s=times,
        temperature_output_c=temperature_output,
        moisture_output=moisture_output,
        temperature_surface_c=temperature_surface,
        moisture_surface=moisture_surface,
        ambient_temperature_c=np.asarray(ambient.temperature(times)),
        ambient_moisture=np.asarray(ambient.concentration(times)),
        heat_boundary_rate_w=heat_boundary_rate,
        heat_storage_cumulative_j=heat_storage_cumulative,
        heat_boundary_cumulative_j=heat_boundary_cumulative,
        moisture_boundary_rate=moisture_boundary_rate,
        moisture_inventory=moisture_inventory,
        cmax=cmax,
        cmax_radius_m=cmax_radius,
        snapshot_times_s=snapshot_times,
        temperature_cells_snapshot_c=snapshot_states[:, :cells],
        moisture_cells_snapshot=snapshot_states[:, cells:],
        terminal_radii_m=terminal_radii,
        terminal_moisture=terminal_moisture,
        end_time_s=end_time_s,
        end_max_moisture=end_max,
        end_max_radius_m=end_radius,
        nfev=int(nfev),
        njev=int(njev),
        nlu=int(nlu),
        runtime_s=elapsed,
        tolerances={
            "rtol": float(rtol),
            "atol_temperature": float(atol_temperature),
            "atol_moisture": float(atol_moisture),
            "max_step_s": float(max_step_s),
        },
    )


def values_at(solution: Q3Solution, times_s: np.ndarray, radii_m: np.ndarray):
    time_indices = np.searchsorted(solution.time_s, times_s)
    if np.any(time_indices >= solution.time_s.size) or not np.allclose(
        solution.time_s[time_indices], times_s, rtol=0.0, atol=1.0e-7
    ):
        raise ValueError("请求的报告时刻不在正式输出中")
    radius_indices = np.searchsorted(OUTPUT_RADII_M, radii_m)
    if not np.allclose(OUTPUT_RADII_M[radius_indices], radii_m, atol=1.0e-12):
        raise ValueError("请求的报告位置不在正式输出中")
    return (
        solution.temperature_output_c[np.ix_(time_indices, radius_indices)],
        solution.moisture_output[np.ix_(time_indices, radius_indices)],
    )


def table5(solution: Q3Solution) -> tuple[np.ndarray, np.ndarray]:
    regular_times = TABLE_TIMES_S[TABLE_TIMES_S < solution.end_time_s - 1.0e-7]
    times = np.r_[regular_times, solution.end_time_s]
    rows = []
    for time_value in times:
        index = int(np.argmin(np.abs(solution.time_s - time_value)))
        if abs(solution.time_s[index] - time_value) > 1.0e-7:
            raise ValueError("表5时刻未包含在解中")
        radius_indices = np.searchsorted(OUTPUT_RADII_M, TABLE_RADII_M)
        rows.append(solution.moisture_output[index, radius_indices])
    return times, np.asarray(rows)


def q2_alignment_check() -> dict[str, object]:
    if not Q2_SOLUTION_PATH.exists():
        raise FileNotFoundError(f"缺少Q2未舍入结果：{Q2_SOLUTION_PATH}")
    alignment_times = np.arange(0.0, 10801.0, 1.0)
    candidate = solve_q3(
        cells=2561,
        max_step_s=10.0,
        integration_upper_s=10800.0,
        terminal_event=False,
        output_times_s=alignment_times,
    )
    fields = (
        "temperature_output_c",
        "moisture_output",
        "temperature_surface_c",
        "moisture_surface",
    )
    report: dict[str, object] = {}
    with np.load(Q2_SOLUTION_PATH) as reference:
        if not np.array_equal(reference["time_s"], alignment_times):
            raise ValueError("Q2未舍入结果的时间轴不是0..10800 s逐秒")
        if not np.array_equal(reference["output_radii_m"], OUTPUT_RADII_M):
            raise ValueError("Q2未舍入结果的21个径向位置与Q3不一致")
        for field in fields:
            current = np.asarray(getattr(candidate, field))
            target = np.asarray(reference[field])
            difference = np.abs(current - target)
            flat = int(np.argmax(difference))
            index = np.unravel_index(flat, difference.shape)
            location: dict[str, float] = {"time_s": float(alignment_times[index[0]])}
            if difference.ndim == 2:
                location["radius_cm"] = float(OUTPUT_RADII_M[index[1]] * 100.0)
            else:
                location["radius_cm"] = 2.0
            report[field] = {
                "max_abs_difference": float(difference[index]),
                "location": location,
            }
    worst_temperature = max(
        report[name]["max_abs_difference"]
        for name in ("temperature_output_c", "temperature_surface_c")
    )
    worst_moisture = max(
        report[name]["max_abs_difference"]
        for name in ("moisture_output", "moisture_surface")
    )
    report["acceptance"] = {
        "temperature_below_1e_8_c": bool(worst_temperature < 1.0e-8),
        "moisture_below_1e_10": bool(worst_moisture < 1.0e-10),
    }
    report["overall_pass"] = bool(all(report["acceptance"].values()))
    report["runtime_s"] = candidate.runtime_s
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    path = RESULTS_DIR / "q3_q2_alignment.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if not report["overall_pass"]:
        raise RuntimeError(f"Q2前3小时对齐失败，详见{path}")
    return report


def save_solution_data(solution: Q3Solution) -> Path:
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    path = RESULTS_DIR / "q3_solution.npz"
    np.savez_compressed(
        path,
        time_s=solution.time_s,
        output_radii_m=OUTPUT_RADII_M,
        temperature_output_c=solution.temperature_output_c,
        moisture_output=solution.moisture_output,
        temperature_surface_c=solution.temperature_surface_c,
        moisture_surface=solution.moisture_surface,
        ambient_temperature_c=solution.ambient_temperature_c,
        ambient_moisture=solution.ambient_moisture,
        heat_boundary_rate_w=solution.heat_boundary_rate_w,
        heat_storage_cumulative_j=solution.heat_storage_cumulative_j,
        heat_boundary_cumulative_j=solution.heat_boundary_cumulative_j,
        moisture_boundary_rate=solution.moisture_boundary_rate,
        moisture_inventory=solution.moisture_inventory,
        cmax=solution.cmax,
        cmax_radius_m=solution.cmax_radius_m,
        snapshot_times_s=solution.snapshot_times_s,
        cell_centers_m=solution.grid.centers_m,
        temperature_cells_snapshot_c=solution.temperature_cells_snapshot_c,
        moisture_cells_snapshot=solution.moisture_cells_snapshot,
        terminal_radii_m=solution.terminal_radii_m,
        terminal_moisture=solution.terminal_moisture,
        end_time_s=np.array(solution.end_time_s),
    )
    return path


def build_result_workbook(
    solution: Q3Solution,
    output_path: Path = DEFAULT_XLSX,
) -> dict[str, object]:
    """复制官方模板，仅写入题目要求的数值。"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.resolve() == OFFICIAL_TEMPLATE.resolve():
        raise ValueError("拒绝覆盖官方result3.xlsx模板")
    shutil.copy2(OFFICIAL_TEMPLATE, output_path)
    workbook = load_workbook(output_path)
    if len(workbook.worksheets) != 1:
        raise ValueError("官方result3.xlsx应仅包含一个工作表")
    worksheet = workbook.worksheets[0]
    template_rows = worksheet.max_row
    template_columns = worksheet.max_column
    regular_mask = (
        (solution.time_s >= OUTPUT_INTERVAL_S - 1.0e-8)
        & (np.abs(solution.time_s / OUTPUT_INTERVAL_S - np.round(solution.time_s / OUTPUT_INTERVAL_S)) < 1.0e-10)
        & (solution.time_s <= solution.end_time_s + 1.0e-8)
    )
    times = solution.time_s[regular_mask]
    values = np.round(solution.moisture_output[regular_mask], 4)
    headers = np.round(OUTPUT_RADII_M * 100.0, 10)
    for column, value in enumerate(headers, start=2):
        source = worksheet.cell(1, min(column, template_columns))
        target = worksheet.cell(1, column)
        target._style = copy.copy(source._style)
        font = copy.copy(source.font)
        font.sz = 10.0
        target.font = font
        target.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.cell(1, column, float(value))
    for row, (time_value, profile) in enumerate(zip(times, values), start=2):
        if row > template_rows:
            worksheet.row_dimensions[row].height = worksheet.row_dimensions[template_rows].height
        source = worksheet.cell(min(row, template_rows), 1)
        target = worksheet.cell(row, 1)
        target._style = copy.copy(source._style)
        font = copy.copy(source.font)
        font.sz = 10.0
        target.font = font
        target.alignment = Alignment(horizontal="center", vertical="center")
        target.value = int(round(time_value))
        for column, value in enumerate(profile, start=2):
            source = worksheet.cell(min(row, template_rows), min(column, template_columns))
            cell = worksheet.cell(row, column)
            cell._style = copy.copy(source._style)
            font = copy.copy(source.font)
            font.sz = 10.0
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.value = float(value)
            cell.number_format = "0.0000"
    for column in range(template_columns + 1, 23):
        letter = worksheet.cell(1, column).column_letter
        source_letter = worksheet.cell(1, template_columns).column_letter
        worksheet.column_dimensions[letter].width = worksheet.column_dimensions[source_letter].width
    workbook.save(output_path)
    return validate_result_workbook(output_path, solution)


def validate_result_workbook(path: Path, solution: Q3Solution | None = None):
    workbook = load_workbook(path, read_only=False, data_only=True)
    template = load_workbook(OFFICIAL_TEMPLATE, read_only=True, data_only=True)
    expected_names = template.sheetnames
    template.close()
    if workbook.sheetnames != expected_names:
        raise ValueError(f"result3.xlsx工作表错误：{workbook.sheetnames}")
    worksheet = workbook.worksheets[0]
    expected_last = int(np.floor((solution.end_time_s if solution else 0.0) / 60.0) * 60)
    expected_rows = expected_last // 60 + 1
    if solution is not None and (worksheet.max_row != expected_rows or worksheet.max_column != 22):
        raise ValueError(
            f"result3.xlsx应为{expected_rows}行×22列，实际为"
            f"{worksheet.max_row}行×{worksheet.max_column}列"
        )
    rows = worksheet.iter_rows(
        min_row=1,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=22,
        values_only=True,
    )
    header_row = next(rows)
    headers = np.asarray(header_row[1:], dtype=float)
    if not np.array_equal(headers, np.round(OUTPUT_RADII_M * 100.0, 10)):
        raise ValueError("result3.xlsx径向表头不一致")
    formula_error_count = 0
    blank_count = 0
    number_format_error_count = 0
    first_time = None
    last_time = None
    for excel_row, row_values in enumerate(rows, start=2):
        expected_time = (excel_row - 1) * 60
        if row_values[0] != expected_time:
            raise ValueError(f"result3.xlsx第{excel_row}行时间错误")
        first_time = expected_time if first_time is None else first_time
        last_time = expected_time
        for column, value in enumerate(row_values[1:], start=2):
            blank_count += int(value is None)
            formula_error_count += int(isinstance(value, str) and value.startswith("#"))
            number_format_error_count += int(
                worksheet.cell(excel_row, column).number_format != "0.0000"
            )
        for column in range(1, 23):
            cell = worksheet.cell(excel_row, column)
            if cell.alignment.horizontal != "center" or cell.alignment.vertical != "center":
                raise ValueError(f"result3.xlsx单元格{cell.coordinate}未居中")
    checks = {
        "sheet_names": workbook.sheetnames,
        "rows": worksheet.max_row,
        "columns": worksheet.max_column,
        "first_time_s": first_time,
        "last_time_s": last_time,
        "first_radius_cm": float(headers[0]),
        "last_radius_cm": float(headers[-1]),
        "blank_data_cells": blank_count,
        "formula_error_count": formula_error_count,
        "number_format_error_count": number_format_error_count,
    }
    workbook.close()
    if blank_count or formula_error_count or number_format_error_count:
        raise ValueError("result3.xlsx存在空洞、公式错误或四位小数格式错误")
    return checks


def save_run_summary(
    solution: Q3Solution,
    data_path: Path,
    workbook_checks: dict[str, object],
    alignment: dict[str, object],
) -> Path:
    table_times, table_values = table5(solution)
    initial = q2.initial_properties()
    summary = {
        "model": "Q3 coupled conservative cell-centered radial finite volume + BDF",
        "grid_cells": solution.grid.cells,
        "method": "BDF",
        "runtime_s": solution.runtime_s,
        "nfev": solution.nfev,
        "njev": solution.njev,
        "nlu": solution.nlu,
        "tolerances": solution.tolerances,
        "input_file": record_path(ATTACHMENT_1),
        "official_template": record_path(OFFICIAL_TEMPLATE),
        "ambient_data_range_s": [0, 14400],
        "ambient_endpoint": {
            "temperature_c": solution.ambient.post_temperature_c,
            "moisture": float(solution.ambient.moisture[-1]),
        },
        "post_14400_environment": {
            "temperature_c": solution.ambient.post_temperature_c,
            "moisture": solution.ambient.post_moisture,
        },
        "initial_properties": initial,
        "threshold": THRESHOLD,
        "end_time_s": solution.end_time_s,
        "end_time_h": solution.end_time_s / 3600.0,
        "end_time_d": solution.end_time_s / 86400.0,
        "end_max_moisture": solution.end_max_moisture,
        "end_max_radius_m": solution.end_max_radius_m,
        "end_axis_moisture": float(solution.moisture_output[-1, 0]),
        "end_surface_moisture": float(solution.moisture_surface[-1]),
        "end_axis_temperature_c": float(solution.temperature_output_c[-1, 0]),
        "end_surface_temperature_c": float(solution.temperature_surface_c[-1]),
        "table5_times_s": table_times.tolist(),
        "table5_radii_cm": (TABLE_RADII_M * 100.0).tolist(),
        "table5_moisture": table_values.tolist(),
        "formal_output": {
            "time_interval_s": 60,
            "last_regular_time_s": int(np.floor(solution.end_time_s / 60.0) * 60),
            "radius_cm": [0.0, 2.0, 0.1],
            "xlsx_decimals": 4,
        },
        "q2_alignment": alignment,
        "solution_data": record_path(data_path),
        "result_workbook": record_path(DEFAULT_XLSX),
        "result_workbook_checks": workbook_checks,
        "all_main_checks_passed": bool(
            alignment["overall_pass"]
            and solution.end_max_moisture <= THRESHOLD + 5.0e-10
            and np.all(np.isfinite(solution.moisture_output))
            and np.min(solution.moisture_output) > 0.0
        ),
    }
    RESULTS_DIR.mkdir(parents=True, exist_ok=True)
    path = RESULTS_DIR / "q3_run_summary.json"
    path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cells", type=int, default=2561)
    parser.add_argument("--rtol", type=float, default=1.0e-8)
    parser.add_argument("--atol-temperature", type=float, default=1.0e-10)
    parser.add_argument("--atol-moisture", type=float, default=1.0e-11)
    parser.add_argument("--max-step", type=float, default=60.0)
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--skip-q2-alignment", action="store_true")
    parser.add_argument("--no-xlsx", action="store_true")
    args = parser.parse_args()

    properties = q2.print_initial_properties()
    ambient = load_q3_ambient()
    print(
        "附件1末端核验："
        f"Ta(14400)={ambient.temperature_c[-1]:.6f} ℃, "
        f"Ca(14400)={ambient.moisture[-1]:.8f} kg/kg"
    )
    alignment = (
        {"overall_pass": True, "skipped_by_user": True}
        if args.skip_q2_alignment
        else q2_alignment_check()
    )
    solution = solve_q3(
        cells=args.cells,
        rtol=args.rtol,
        atol_temperature=args.atol_temperature,
        atol_moisture=args.atol_moisture,
        max_step_s=args.max_step,
    )
    data_path = save_solution_data(solution)
    workbook_checks: dict[str, object] = {"not_generated": True}
    if not args.no_xlsx:
        workbook_checks = build_result_workbook(solution, args.xlsx)
    summary_path = save_run_summary(solution, data_path, workbook_checks, alignment)
    table_times, table_values = table5(solution)
    print(f"初始D={properties['D_m2_s']:.10e} m^2/s")
    print(
        f"Q3正式计算完成：N={args.cells}, t_end={solution.end_time_s:.6f} s "
        f"({solution.end_time_s/3600.0:.9f} h), runtime={solution.runtime_s:.3f} s"
    )
    print(
        f"终点Cmax={solution.end_max_moisture:.12f}, "
        f"位置r={solution.end_max_radius_m*100.0:.9f} cm"
    )
    print("表5时刻（h）：", table_times / 3600.0)
    print(np.array2string(table_values, precision=10, suppress_small=False))
    print(f"未舍入结果：{data_path}")
    print(f"运行摘要：{summary_path}")
    if not args.no_xlsx:
        print(f"结果工作簿：{args.xlsx}")
        print(json.dumps(workbook_checks, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
